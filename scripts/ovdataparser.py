import argparse
import json
import re
from pathlib import Path
from typing import NotRequired, TypedDict

from openpyxl import load_workbook


class BaderData(TypedDict):
    vacancy_nearest_neighbors: NotRequired[float]
    by_element: NotRequired[dict[str, float]]
    per_atom: NotRequired[dict[str, float]]
    total_per_atom: NotRequired[float]
    volume: NotRequired[float]


class VolumeData(TypedDict):
    vacancy: NotRequired[float]
    perfect: NotRequired[float]
    difference: NotRequired[float]
    ratio: NotRequired[float]


class DisplacementData(TypedDict):
    initial: NotRequired[float]
    final: NotRequired[float]
    difference: NotRequired[float]


class VacancyRecord(TypedDict):
    system: str
    index: NotRequired[int]
    vacancy_position: int

    defect_energy: NotRequired[float]
    perfect_energy: NotRequired[float]
    oxygen_vacancy_energy: NotRequired[float]
    vacancy_formation_energy: NotRequired[float]

    neighbor_composition: NotRequired[dict[str, int]]
    bader: NotRequired[BaderData]
    volume: NotRequired[VolumeData]
    displacement: NotRequired[DisplacementData]


SYSTEM_NAME_MAP = {
    "MgO": "MgO",
    "MgCuO+U2": "MgCuO",
    "MgZnO+U": "MgZnO",
    "MgCo+U": "MgCoO",
    "MgNi+U": "MgNiO",
    "5HEO+U": "MgCoCuNiZnO",
    "MgNiZnO": "MgNiZnO",
    "MgNiCuO": "MgNiCuO",
    "MgCuZnO": "MgCuZnO",
    "MgCoZnO": "MgCoZnO",
    "MgCoNiO": "MgCoNiO",
    "MgCoCuO": "MgCoCuO",
    "MgCuNiZnO": "MgCuNiZnO",
    "MgCuCoZnO": "MgCuCoZnO",
    "MgCoNiZnO": "MgCoNiZnO",
    "MgCoCuNiO": "MgCoCuNiO",
    "MgCoCuNiZnO": "MgCoCuNiZnO",
}


def normalize_system_name(system_name: str) -> str:
    return SYSTEM_NAME_MAP.get(system_name, system_name)


def parse_vacancy_position(value: object) -> int:
    match = re.fullmatch(r"Atom(\d+)", str(value).strip(), re.IGNORECASE)
    if not match:
        raise ValueError(f"Invalid vacancy position: {value!r}")
    return int(match.group(1))


def convert_sheet(ws) -> list[VacancyRecord]:
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(ws[1])
        if cell.value is not None
    }

    if "OV Posit" not in headers:
        return []

    records: list[VacancyRecord] = []

    fixed_fields = {
        "Defect(eV)": "defect_energy",
        "Perfect(eV)": "perfect_energy",
        "OV (eV)": "oxygen_vacancy_energy",
        "OE(eV)": "oxygen_vacancy_energy",
        "V-Formation(eV)": "vacancy_formation_energy",
        "OVF(eV)": "vacancy_formation_energy",
    }

    for values in ws.iter_rows(min_row=2, values_only=True):
        vacancy_value = values[headers["OV Posit"]]
        if vacancy_value is None:
            continue

        record: VacancyRecord = {
            "system": normalize_system_name(ws.title),
            "vacancy_position": parse_vacancy_position(vacancy_value),
        }

        if "No" in headers:
            value = values[headers["No"]]
            if value is not None:
                record["index"] = int(value)

        for excel_name, json_name in fixed_fields.items():
            if excel_name not in headers:
                continue

            value = values[headers[excel_name]]
            if value is not None:
                record[json_name] = float(value)

        neighbor_composition: dict[str, int] = {}
        bader_by_element: dict[str, float] = {}
        bader_per_atom: dict[str, float] = {}

        for header, column_index in headers.items():
            value = values[column_index]
            if value is None:
                continue

            if header.startswith("No of "):
                element = header.removeprefix("No of ").strip()
                neighbor_composition[element] = int(value)

            elif (
                header.startswith("Bader_")
                and not header.startswith("Bader_per")
                and header not in {"Bader_OV_NN", "Bader_vol"}
            ):
                element = header.removeprefix("Bader_").strip()
                bader_by_element[element] = float(value)

            elif header.startswith("Bader_per"):
                element = header.removeprefix("Bader_per").strip()
                bader_per_atom[element] = float(value)

        if neighbor_composition:
            record["neighbor_composition"] = neighbor_composition

        bader: BaderData = {}

        if "Bader_OV_NN" in headers:
            value = values[headers["Bader_OV_NN"]]
            if value is not None:
                bader["vacancy_nearest_neighbors"] = float(value)

        if bader_by_element:
            bader["by_element"] = bader_by_element

        if bader_per_atom:
            bader["per_atom"] = bader_per_atom

        if "TBC/atom" in headers:
            value = values[headers["TBC/atom"]]
            if value is not None:
                bader["total_per_atom"] = float(value)

        if "Bader_vol" in headers:
            value = values[headers["Bader_vol"]]
            if value is not None:
                bader["volume"] = float(value)

        if bader:
            record["bader"] = bader

        volume: VolumeData = {}

        for header in ("V-O volume(Å)", "V_vac"):
            if header not in headers:
                continue

            value = values[headers[header]]
            if value is not None:
                volume["vacancy"] = float(value)
                break

        for header in ("O_vol_pef", "V-perf"):
            if header not in headers:
                continue

            value = values[headers[header]]
            if value is not None:
                volume["perfect"] = float(value)
                break

        if "vol_diff" in headers:
            value = values[headers["vol_diff"]]
            if value is not None:
                volume["difference"] = float(value)

        # Prefer vv/vp when both columns are present because it has greater
        # precision. Fall back to dV only when vv/vp is absent or empty.
        ratio = None

        if "vv/vp" in headers:
            value = values[headers["vv/vp"]]
            if value is not None:
                ratio = float(value)

        if ratio is None and "dV" in headers:
            value = values[headers["dV"]]
            if value is not None:
                ratio = float(value)

        if ratio is not None:
            volume["ratio"] = ratio

        if volume:
            record["volume"] = volume

        displacement: DisplacementData = {}
        displacement_fields = {
            "dl_i": "initial",
            "dl_f": "final",
            "Dd": "difference",
        }

        for excel_name, json_name in displacement_fields.items():
            if excel_name not in headers:
                continue

            value = values[headers[excel_name]]
            if value is not None:
                displacement[json_name] = float(value)

        if displacement:
            record["displacement"] = displacement

        records.append(record)

    return records


def convert_workbook(input_path: Path, output_path: Path) -> int:
    workbook = load_workbook(input_path, data_only=True)

    records: list[VacancyRecord] = []

    for ws in workbook.worksheets:
        records.extend(convert_sheet(ws))

    with output_path.open("w", encoding="ascii") as file:
        json.dump(records, file, indent=2, allow_nan=False)

    return len(records)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert oxygen-vacancy Excel data to JSON."
    )
    parser.add_argument(
        "input",
        type=Path,
        help="Input Excel workbook.",
    )
    parser.add_argument(
        "output",
        type=Path,
        help="Output JSON file.",
    )
    args = parser.parse_args()

    record_count = convert_workbook(args.input, args.output)
    print(f"Wrote {record_count} records to {args.output}")


if __name__ == "__main__":
    main()
